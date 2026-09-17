import os
import logging
from pathlib import Path
from typing import Dict, Any, Optional
from fpdf import FPDF
import docx
from backend.config import settings
from backend.services.llm_service import LLMService

logger = logging.getLogger(__name__)


class ReportAgent:
    @staticmethod
    def generate_report(
        title: str,
        query: str,
        final_response: str,
        citations: list,
        format_type: str = "markdown"  # markdown, pdf, docx
    ) -> Dict[str, Any]:
        logger.info(f"Running Report Agent to produce {format_type} report: '{title}'")

        prompt = f"""Generate a comprehensive formal report document.

Title: {title}
Original Request: {query}

Verified Findings & Content:
{final_response}

Instructions:
Include the following structured sections:
1. Executive Summary
2. Key Findings & Requirements
3. Detailed Evidence Analysis
4. Recommendations & Action Items
5. References & Document Sources
"""
        report_markdown = LLMService.generate(prompt)
        if not report_markdown or report_markdown.startswith("[Simulated"):
            report_markdown = f"# {title}\n\n## Executive Summary\n{final_response}\n\n## References\n" + "\n".join([f"- {c.get('document')} Page {c.get('page')}" for c in citations])

        report_dir = Path(settings.REPORTS_DIR)
        report_dir.mkdir(parents=True, exist_ok=True)
        filename_base = title.lower().replace(" ", "_").replace("/", "_")[:40]

        files_created = {}

        # 1. Save Markdown
        md_filename = f"{filename_base}.md"
        md_path = report_dir / md_filename
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(report_markdown)
        files_created["markdown"] = str(md_path)

        # 2. Save DOCX
        try:
            docx_filename = f"{filename_base}.docx"
            docx_path = report_dir / docx_filename
            doc = docx.Document()
            doc.add_heading(title, 0)
            for line in report_markdown.split("\n"):
                if line.startswith("# "):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith("## "):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith("### "):
                    doc.add_heading(line[4:], level=3)
                elif line.strip():
                    doc.add_paragraph(line)
            doc.save(docx_path)
            files_created["docx"] = str(docx_path)
        except Exception as e:
            logger.warning(f"Could not generate DOCX report: {e}")

        # 3. Save PDF
        try:
            pdf_filename = f"{filename_base}.pdf"
            pdf_path = report_dir / pdf_filename
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Helvetica", size=12)
            # FPDF clean text
            clean_text = report_markdown.encode('latin-1', 'replace').decode('latin-1')
            for line in clean_text.split("\n"):
                if line.strip():
                    pdf.multi_cell(0, 8, text=line)
                    pdf.ln(2)
            pdf.output(str(pdf_path))
            files_created["pdf"] = str(pdf_path)
        except Exception as e:
            logger.warning(f"Could not generate PDF report: {e}")

        return {
            "title": title,
            "report_markdown": report_markdown,
            "files": files_created,
            "primary_file": files_created.get(format_type, files_created.get("markdown"))
        }
